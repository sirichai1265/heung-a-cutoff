#!/usr/bin/env python3
"""Save the data behind index.html as an English Excel snapshot in daily/<date>/.

File name carries the run date and time (Bangkok): daily/2026-10-11/Cut off and 1st return - 2026-10-11 1054.xlsx
Run after build_dashboard.py. Needs openpyxl.
"""
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).parent
COLS = ["POL", "Service", "Vessel", "Vessel Name", "Voyage", "Wharf", "Wharf Name", "Next Port",
        "1st Return", "Cut off (Dry)", "Cut off (Reefer)", "ETA", "ETB", "ETD", "ETA Source"]
WIDTHS = [8, 9, 9, 22, 9, 8, 44, 10, 13, 18, 18, 18, 18, 18, 30]
SRC = {"web": "Estimated from site berth time", "xls": "Internal CUT file", "est": "Estimated"}


def main():
    html = (HERE / "index.html").read_text(encoding="utf-8")
    raw = json.loads(re.search(r"const RAW = (\[.*?\]);\nconst WHARF_MAP", html, re.S).group(1))
    wm = json.loads(re.search(r"const WHARF_MAP = (\{.*?\});", html, re.S).group(1))
    P = lambda s: datetime.strptime(s, "%Y-%m-%d %H:%M")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cut off and 1st return"
    ws.append(COLS)
    for r in sorted(raw, key=lambda r: (r["pol"], r["eta"])):
        ws.append([r["pol"], r["service"], r["vessel_code"], r["vessel"], r["vyg_bound"], r["wharf"],
                   wm.get(r["wharf"], ""), r["pod"], P(r["opengate"]), P(r["cutoff_dry"]),
                   P(r["cutoff_reefer"]), P(r["eta"]), P(r["etb"]), P(r["etd"]), SRC[r["eta_src"]]])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B2540")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        row[8].number_format = "[$-409]dd mmm yyyy"  # English dates whatever the PC locale
        for i in range(9, 14):
            row[i].number_format = "[$-409]dd mmm yyyy hh:mm"
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    log_file = HERE / "history" / "changes.json"
    if log_file.exists():  # day-by-day change log (see history.py)
        cl = wb.create_sheet("Change log")
        cl.append(["Date", "Compared with", "POL", "Service", "Vessel", "Vessel Name", "Voyage", "Change", "Change (h)",
                   "ETB (old)", "ETB (new)", "ETD (old)", "ETD (new)", "Cut off Dry (old)", "Cut off Dry (new)",
                   "Cut off Reefer (old)", "Cut off Reefer (new)"])
        days = json.loads(log_file.read_text(encoding="utf-8"))["days"]
        for day in sorted(days, reverse=True):
            for e in days[day]["events"]:
                o, n = e["old"], e.get("new") or {}
                pt = lambda d, k: P(d[k]) if d.get(k) else None
                cl.append([datetime.strptime(day, "%Y-%m-%d"), datetime.strptime(days[day]["vs"], "%Y-%m-%d"), e["pol"], e["service"],
                           e["vessel_code"], e["vessel"], e["vyg_bound"], "Delayed" if e["type"] == "changed" and e["chg_h"] > 0 else
                           ("Advanced" if e["type"] == "changed" else "Removed"), e.get("chg_h"),
                           pt(o, "etb"), pt(n, "etb"), pt(o, "etd"), pt(n, "etd"), pt(o, "dry"), pt(n, "dry"), pt(o, "rf"), pt(n, "rf")])
        for c in cl[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B2540")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in cl.iter_rows(min_row=2):
            row[0].number_format = row[1].number_format = "[$-409]dd mmm yyyy"
            for i in range(9, 17):
                row[i].number_format = "[$-409]dd mmm yyyy hh:mm"
        for i, w in enumerate([13, 13, 8, 9, 9, 22, 9, 10, 10] + [18] * 8):
            cl.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        cl.freeze_panes = "A2"
        cl.auto_filter.ref = cl.dimensions

    now = datetime.now(timezone(timedelta(hours=7)))
    out = HERE / "daily" / f"{now:%Y-%m-%d}" / f"Cut off and 1st return - {now:%Y-%m-%d %H%M}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Saved {out.relative_to(HERE)} ({len(raw)} voyages)")


if __name__ == "__main__":
    main()
